import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal

from openpyxl import load_workbook

from qlens.errors import ApiError

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 60
_DELIMITERS = ",;\t|"


@dataclass
class RawTable:
    labels: list[str]               # headers as written (stripped; blank ones become column_<n>)
    names: list[str]                # sanitized identifiers, same order
    rows: list[list[str | None]]


def _fail(status: int, code: str, message: str) -> ApiError:
    return ApiError(status, code, message)


# ---- header names ---------------------------------------------------------------------------------

def sanitize_names(labels: list[str]) -> list[str]:
    taken: set[str] = set()
    out: list[str] = []
    for position, label in enumerate(labels, start=1):
        folded = unicodedata.normalize("NFKD", label).encode("ascii", "ignore").decode("ascii")
        name = re.sub(r"[^a-z0-9]+", "_", folded.lower()).strip("_")
        if not name:
            name = f"column_{position}"
        elif name[0].isdigit():
            name = f"c_{name}"
        name = name[:63].rstrip("_")
        candidate, n = name, 1
        while candidate in taken:
            n += 1
            suffix = f"_{n}"
            candidate = name[: 63 - len(suffix)].rstrip("_") + suffix
        taken.add(candidate)
        out.append(candidate)
    return out


# ---- reading --------------------------------------------------------------------------------------

def read_table(filename: str, data: bytes) -> RawTable:
    if len(data) > MAX_BYTES:
        raise _fail(422, "file_too_large", "The file is larger than 20 MB.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("csv", "tsv", "txt"):
        grid = _csv_rows(data)
    elif ext == "xlsx":
        grid = _xlsx_rows(data)
    else:
        raise _fail(422, "unsupported_file", "Upload a .csv or .xlsx file.")
    return _build_table(grid)


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows saves "CSV" as cp1252; bytes cp1252 leaves undefined become U+FFFD.
        return data.decode("cp1252", errors="replace")


def _sniff_delimiter(text: str) -> str:
    sample = text[:65536]
    first_line = next((ln for ln in sample.splitlines() if ln.strip()), "")
    try:
        found = csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        return ","
    # The sniffer can latch onto a character that is only frequent in the data; a real
    # delimiter always appears in the header line.
    return found if found in first_line else ","


def _csv_rows(data: bytes) -> Iterator[list[str | None]]:
    text = _decode(data)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text))
    try:
        for row in reader:
            yield [c.replace("\x00", "").strip() or None for c in row]   # NUL cannot be stored in text columns
    except csv.Error as exc:
        raise _fail(422, "unsupported_file", f"Could not read the CSV file: {exc}") from exc


def _xlsx_cell(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat()
    if isinstance(v, int):
        return str(v)
    if isinstance(v, (float, Decimal)):
        d = Decimal(repr(v)) if isinstance(v, float) else v
        if d != d.to_integral_value():
            return format(d, "f")
        return str(int(d))              # 4000.0 -> "4000"
    s = str(v).replace("\x00", "").strip()
    return s or None


def _xlsx_rows(data: bytes) -> Iterator[list[str | None]]:
    bad = _fail(422, "unsupported_file", "This does not look like a valid .xlsx file.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several of its own types for corrupt packages
        raise bad from exc
    try:
        if not wb.worksheets:
            return
        try:
            for row in wb.worksheets[0].iter_rows(values_only=True):
                yield [_xlsx_cell(v) for v in row]
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise bad from exc
    finally:
        wb.close()


def _build_table(grid: Iterable[list[str | None]]) -> RawTable:
    header: list[str | None] | None = None
    body: list[list[str | None]] = []
    for row in grid:
        if not any(c is not None for c in row):
            continue                                    # fully empty rows are skipped
        if header is None:
            header = row
            if len(header) > MAX_COLUMNS and any(c is not None for c in header[MAX_COLUMNS:]):
                raise _fail(422, "too_many_columns", f"At most {MAX_COLUMNS} columns are supported.")
            continue
        if len(body) >= MAX_ROWS:
            raise _fail(422, "too_many_rows", f"At most {MAX_ROWS:,} data rows are supported.")
        body.append(row)
    if header is None or not body:
        raise _fail(422, "empty_file", "The file has no data rows.")

    # Width = the header's width minus trailing columns that are blank in the header and empty in
    # every row (spreadsheets often carry hundreds of formatted-but-empty columns).
    width = len(header)
    while width > 0 and header[width - 1] is None and all(
            len(r) < width or r[width - 1] is None for r in body):
        width -= 1
    if width == 0:
        raise _fail(422, "empty_file", "The file has no data rows.")
    if width > MAX_COLUMNS:
        raise _fail(422, "too_many_columns", f"At most {MAX_COLUMNS} columns are supported.")

    names = sanitize_names([h or "" for h in header[:width]])
    labels = [h if h is not None else f"column_{i}" for i, h in enumerate(header[:width], start=1)]
    rows = [(r + [None] * width)[:width] for r in body]
    return RawTable(labels=labels, names=names, rows=rows)
